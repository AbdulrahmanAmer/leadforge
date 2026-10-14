"""Import hand-checked outreach outcomes for enrichment benchmarking."""

from __future__ import annotations

import csv
import json
import re
import sqlite3
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from leadforge import db
from leadforge.models import Contact, EnrichmentObservation, Person
from leadforge.normalize import to_e164
from leadforge.util import now_iso


def _norm(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _phone(value: Any) -> str:
    return re.sub(r"[^+0-9]", "", str(value or ""))


def _domain(value: Any) -> str:
    """Return a conservative hostname key for workbook/CSV matching."""
    text = str(value or "").strip().casefold()
    if not text:
        return ""
    parsed = urlparse(text if "://" in text else f"https://{text}")
    host = (parsed.hostname or "").removeprefix("www.")
    return host


def _canonical_outcome(value: Any) -> str:
    raw = str(value or "").strip().casefold()
    return {
        "meeting booked": "meeting_booked",
        "interested / qualified": "interested",
        "interested": "interested",
        "callback": "callback",
        "send information": "send_information",
        "voicemail": "voicemail",
        "no answer": "no_answer",
        "wrong / disconnected": "wrong_number",
        "not interested": "not_interested",
        "do not contact": "do_not_contact",
        "dnc": "do_not_contact",
    }.get(raw, raw.replace("/", "_").replace(" ", "_") or "unknown")


_WORKBOOK_HEADER_ALIASES = {
    "queue": "queue",
    "queuenumber": "queue",
    "company": "company",
    "decisionmaker": "decision_maker",
    "decisionmakeraskfor": "decision_maker",
    "dmname": "decision_maker",
    "dmemail": "email",
    "email": "email",
    "phone": "phone",
    "primaryphone": "phone",
    "dmdirectphonebestroute": "dm_phone",
    "dmtitle": "dm_title",
    "website": "website",
    "city": "city",
    "disposition": "outcome",
    "outcome": "outcome",
    "notes": "notes",
    "dialedby": "recorded_by",
    "agent": "recorded_by",
    "dialedat": "contacted_at",
    "timestampcairo": "contacted_at",
    "originalcallbacknote": "callback_note",
    "opportunitysignal": "opportunity_signal",
    "priority": "priority",
    "source": "source",
    "provider": "provider",
    "contacttype": "contact_type",
    "campaign": "campaign",
}


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _workbook_rows(path: Path) -> list[dict[str, str]]:
    """Read the known WESCORE calling-workbook layouts without editing them."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is in the core wheel
        raise ValueError("XLSX benchmark import requires openpyxl") from exc

    rows: list[dict[str, str]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    preferred = ["Dialing List", "List #2", "List #3", "Call Log", "Habiba Callbacks", "Nagham Callbacks"]
    ordered = [name for name in preferred if name in workbook.sheetnames]
    ordered.extend(name for name in workbook.sheetnames if name not in ordered and name != "Dashboard")
    for sheet_name in ordered:
        sheet = workbook[sheet_name]
        materialized = list(sheet.iter_rows(values_only=True))
        header_index = next(
            (i for i, values in enumerate(materialized[:6])
             if any(_header_key(value) in {"company", "queuenumber", "phone", "primaryphone"} for value in values)),
            None,
        )
        if header_index is None:
            continue
        headers = [_WORKBOOK_HEADER_ALIASES.get(_header_key(value), "") for value in materialized[header_index]]
        for values in materialized[header_index + 1:]:
            row: dict[str, str] = {"sheet": sheet_name}
            for key, value in zip(headers, values, strict=False):
                if key and value not in (None, ""):
                    row[key] = str(value).strip()
            if row.get("company") or row.get("phone") or row.get("email"):
                if sheet_name.endswith("Callbacks"):
                    row["outcome"] = "callback"
                    row["notes"] = " | ".join(
                        part for part in (row.get("callback_note"), row.get("opportunity_signal"), row.get("notes")) if part
                    )
                rows.append(row)
    return rows


def _input_rows(path: Path) -> list[dict[str, str]]:
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        return _workbook_rows(path)
    with path.open(newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(2048)
        handle.seek(0)
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t")
        return [
            {str(k).casefold().strip(): str(v or "").strip() for k, v in raw.items()}
            for raw in csv.DictReader(handle, dialect=dialect)
        ]


def _match_business(conn: sqlite3.Connection, row: dict[str, str]) -> tuple[str | None, str]:
    business_id = row.get("business_id", "").strip()
    if business_id and conn.execute("SELECT 1 FROM businesses WHERE id=?", (business_id,)).fetchone():
        return business_id, "business_id"

    domain = _domain(row.get("domain") or row.get("website"))
    if domain:
        hits = conn.execute(
            "SELECT id FROM businesses WHERE lower(replace(domain,'www.',''))=?", (domain,)
        ).fetchall()
        if len(hits) == 1:
            return hits[0]["id"], "domain"
        if len(hits) > 1:
            return None, "ambiguous_domain"

    phone = _phone(row.get("phone"))
    if phone:
        hits = conn.execute("SELECT id FROM businesses WHERE phone_e164=?", (phone,)).fetchall()
        if len(hits) == 1:
            return hits[0]["id"], "phone"
        if len(hits) > 1:
            return None, "ambiguous_phone"

    email = row.get("email", "").strip().casefold()
    if email:
        hits = conn.execute(
            "SELECT DISTINCT business_id FROM contacts WHERE lower(value)=?", (email,)
        ).fetchall()
        if len(hits) == 1:
            return hits[0]["business_id"], "contact"
        if len(hits) > 1:
            return None, "ambiguous_contact"
    return None, "unmatched"


def import_outcomes(conn: sqlite3.Connection, path: Path, dry_run: bool = True) -> dict[str, Any]:
    """Preview or import hand-checked outcomes from CSV/TSV/XLSX using strong matching only."""
    result: dict[str, Any] = {
        "rows": 0, "matched": 0, "imported": 0, "duplicates": 0,
        "ambiguous": 0, "unmatched": 0, "reasons": {}
    }
    for row in _input_rows(path):
        result["rows"] += 1
        business_id, reason = _match_business(conn, row)
        result["reasons"][reason] = result["reasons"].get(reason, 0) + 1
        if not business_id:
            result["ambiguous" if reason.startswith("ambiguous") else "unmatched"] += 1
            continue
        result["matched"] += 1
        if dry_run:
            continue
        channel = row.get("channel") or "phone"
        outcome = _canonical_outcome(row.get("outcome"))
        contacted_at = row.get("call_date") or row.get("contacted_at") or None
        notes = row.get("notes", "")
        duplicate = conn.execute(
            "SELECT 1 FROM outcomes WHERE business_id=? AND channel=? AND result=? "
            "AND (? IS NULL OR contacted_at=?) AND COALESCE(notes,'')=? LIMIT 1",
            (business_id, channel, outcome, contacted_at, contacted_at, notes),
        ).fetchone()
        if duplicate:
            result["duplicates"] += 1
            continue
        db.add_outcome(
            conn, business_id, channel, outcome,
            campaign=row.get("campaign", "") or row.get("sheet", ""), notes=notes,
            recorded_by=row.get("caller") or row.get("recorded_by", ""),
            contacted_at=contacted_at,
            source=row.get("source") or f"user_workbook:{row.get('sheet', 'unknown')}",
            provider=row.get("provider") or "human_calling",
            contact_type=row.get("contact_type") or ("dm_direct" if row.get("dm_phone") else "phone"),
        )
        result["imported"] += 1
    return result


def import_dm_fields(conn: sqlite3.Connection, path: Path, dry_run: bool = True,
                     default_region: str = "US") -> dict[str, Any]:
    """Import operator-confirmed DM fields as a calibration set.

    This is deliberately separate from ``import_outcomes``: a calling workbook is not public-source
    evidence. It is human-observed BYOL data, retained with sheet provenance and never allowed to
    overwrite an OSS observation. Matching still requires a unique business match.
    """
    result: dict[str, Any] = {
        "rows": 0, "matched": 0, "imported": 0, "duplicates": 0,
        "ambiguous": 0, "unmatched": 0, "reasons": {},
    }
    for row in _input_rows(path):
        result["rows"] += 1
        match_row = {**row, "phone": row.get("phone") or row.get("dm_phone")}
        business_id, reason = _match_business(conn, match_row)
        result["reasons"][reason] = result["reasons"].get(reason, 0) + 1
        if not business_id:
            result["ambiguous" if reason.startswith("ambiguous") else "unmatched"] += 1
            continue
        result["matched"] += 1
        if dry_run:
            continue
        name = row.get("decision_maker", "").strip()
        title = row.get("dm_title", "").strip()
        dm_phone = row.get("dm_phone", "").strip()
        email = row.get("email", "").strip().casefold()
        if not name and not dm_phone and not email:
            continue
        person_id = None
        if name:
            db.add_person(conn, Person(business_id=business_id, name=name, title=title,
                                       labeled_by="byol", origin="human_calling", labeled_at=now_iso()))
            person_row = conn.execute(
                "SELECT id FROM people WHERE business_id=? AND name=? AND COALESCE(title,'')=?",
                (business_id, name, title),
            ).fetchone()
            person_id = person_row["id"] if person_row else None
        source = f"human_calling:{row.get('sheet', 'unknown')}"
        source_url = row.get("website", "")
        evidence = f"Operator workbook {row.get('sheet', 'unknown')}: {name} — {title}".strip(" —")
        if dm_phone:
            normalized, _reason = to_e164(dm_phone, default_region)
            value = normalized or _phone(dm_phone)
            duplicate = conn.execute(
                "SELECT 1 FROM enrichment_observations WHERE business_id=? AND field='phone' AND value_json=? "
                "AND source=? AND COALESCE(person_id,0)=COALESCE(?,0) LIMIT 1",
                (business_id, json.dumps(value), source, person_id),
            ).fetchone()
            if duplicate:
                result["duplicates"] += 1
            else:
                db.add_contact(conn, Contact(business_id=business_id, kind="phone", value=value,
                                             label="human_called_dm", tier="unknown", phone_type="direct"))
                db.add_observation(conn, EnrichmentObservation(
                    business_id=business_id, field="phone", value=value, source=source,
                    source_url=source_url, evidence=evidence, confidence=1.0,
                    verification_state="verified", provider="human_calling", person_id=person_id,
                    person_name=name, observed_at=now_iso(), jurisdiction=default_region,
                ))
                result["imported"] += 1
        if email:
            duplicate = conn.execute(
                "SELECT 1 FROM enrichment_observations WHERE business_id=? AND field='email' AND value_json=? "
                "AND source=? AND COALESCE(person_id,0)=COALESCE(?,0) LIMIT 1",
                (business_id, json.dumps(email), source, person_id),
            ).fetchone()
            if duplicate:
                result["duplicates"] += 1
            else:
                db.add_contact(conn, Contact(business_id=business_id, kind="email", value=email,
                                             label="human_called_dm", tier="unknown", email_stage="unknown"))
                db.add_observation(conn, EnrichmentObservation(
                    business_id=business_id, field="email", value=email, source=source,
                    source_url=source_url, evidence=evidence, confidence=1.0,
                    verification_state="observed", provider="human_calling", person_id=person_id,
                    person_name=name, observed_at=now_iso(), jurisdiction=default_region,
                ))
                result["imported"] += 1
    if not dry_run:
        conn.commit()
    return result
