import json
from pathlib import Path

from openpyxl import Workbook
from typer.testing import CliRunner

from leadforge import db
from leadforge.cli import app
from leadforge.config import load_config
from leadforge.enrich.benchmark import import_dm_fields, import_outcomes
from leadforge.models import Business


def test_outcome_benchmark_import_is_dry_run_and_strong_match_only(conn, tmp_path):
    db.upsert_business(conn, Business(id="b1", name="Acme", name_norm="acme", domain="acme.test", dedupe_key="k1"))
    path = tmp_path / "called.tsv"
    path.write_text(
        "business_id\tchannel\toutcome\tcall_date\tcontact_type\tsource\n"
        "b1\tphone\tinterested\t2026-10-14\tmain_line\tgbp\n"
        "unknown\tphone\tno_answer\t2026-10-14\tmain_line\tgbp\n", encoding="utf-8"
    )
    preview = import_outcomes(conn, path)
    assert preview["matched"] == 1 and preview["imported"] == 0
    assert db.outcome_counts(conn) == {}
    written = import_outcomes(conn, path, dry_run=False)
    assert written["imported"] == 1
    row = conn.execute("SELECT * FROM outcomes WHERE business_id='b1'").fetchone()
    assert row["source"] == "gbp" and row["contact_type"] == "main_line"


def test_benchmark_outcomes_cli_is_preview_by_default(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    cfg = load_config(tmp_path)
    conn = db.connect(cfg.db_path)
    db.upsert_business(conn, Business(id="b1", name="Acme", name_norm="acme", dedupe_key="k1"))
    conn.close()
    path = tmp_path / "called.csv"
    path.write_text("business_id,channel,outcome\nb1,phone,interested\n", encoding="utf-8")
    result = CliRunner().invoke(app, ["--json", "benchmark", "outcomes", "--in", str(path)])
    assert result.exit_code == 0, result.output
    digest = next(line for line in result.output.splitlines() if line.startswith("LF_DIGEST "))
    assert json.loads(digest[len("LF_DIGEST "):])["counts"]["imported"] == 0


def test_workbook_import_reads_calling_tabs_and_is_idempotent(conn, tmp_path: Path):
    business_id = "workbook-business"
    db.upsert_business(conn, Business(id=business_id, name="Example", name_norm="example", domain="example.co.uk", dedupe_key="workbook"))
    conn.execute("UPDATE businesses SET domain='example.co.uk', phone_e164='+442012345678' WHERE id=?", (business_id,))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dialing List"
    sheet.append(["Queue #", "Company", "Decision Maker", "Phone", "Disposition", "Notes", "Website", "Dialed By"])
    sheet.append(["1", "Example", "Jane Doe", "+44 20 1234 5678", "Callback", "tomorrow", "https://www.example.co.uk", "Agent"])
    callbacks = workbook.create_sheet("Nagham Callbacks")
    callbacks.append(["NAGHAM CALLBACKS"])
    callbacks.append(["", ""])
    callbacks.append(["Priority", "Queue #", "Company", "Phone", "Original Callback Note", "Website"])
    callbacks.append(["A", "1", "Example", "+44 20 1234 5678", "tomorrow", "https://www.example.co.uk"])
    path = tmp_path / "called.xlsx"
    workbook.save(path)

    preview = import_outcomes(conn, path)
    assert preview["rows"] == 2
    assert preview["matched"] == 2
    written = import_outcomes(conn, path, dry_run=False)
    assert written["imported"] == 1
    assert written["duplicates"] == 1
    assert conn.execute("SELECT COUNT(*) FROM outcomes WHERE business_id=?", (business_id,)).fetchone()[0] == 1


def test_workbook_dm_fields_are_person_owned_and_dry_run_by_default(conn, tmp_path: Path):
    db.upsert_business(conn, Business(id="b1", name="Example", name_norm="example",
                                      domain="example.co.uk", dedupe_key="dm-fields"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "List #2"
    sheet.append(["Company", "Decision Maker", "DM Title", "DM Direct Phone / Best Route", "DM Email", "Website"])
    sheet.append(["Example", "Jane Doe", "Owner", "+44 20 1234 5678", "jane@example.co.uk",
                  "https://www.example.co.uk"])
    path = tmp_path / "dm-fields.xlsx"
    workbook.save(path)

    preview = import_dm_fields(conn, path, default_region="GB")
    assert preview["matched"] == 1 and preview["imported"] == 0
    assert conn.execute("SELECT COUNT(*) FROM people").fetchone()[0] == 0

    written = import_dm_fields(conn, path, dry_run=False, default_region="GB")
    assert written["imported"] == 2
    person = conn.execute("SELECT * FROM people WHERE name='Jane Doe'").fetchone()
    assert person["title"] == "Owner"
    observations = conn.execute(
        "SELECT * FROM enrichment_observations WHERE provider='human_calling' ORDER BY field"
    ).fetchall()
    assert {row["field"] for row in observations} == {"email", "phone"}
    assert all(row["person_id"] == person["id"] for row in observations)
    assert all(row["source"] == "human_calling:List #2" for row in observations)


def test_dm_route_outcomes_are_attributed_as_dm_direct(conn, tmp_path: Path):
    db.upsert_business(conn, Business(id="b1", name="Example", name_norm="example",
                                      domain="example.co.uk", dedupe_key="dm-outcome"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "List #2"
    sheet.append(["Company", "Decision Maker", "DM Direct Phone / Best Route", "Disposition", "Website"])
    sheet.append(["Example", "Jane Doe", "+44 20 1234 5678", "Callback", "https://www.example.co.uk"])
    path = tmp_path / "dm-outcome.xlsx"
    workbook.save(path)

    result = import_outcomes(conn, path, dry_run=False)
    assert result["imported"] == 1
    outcome = conn.execute("SELECT * FROM outcomes WHERE business_id='b1'").fetchone()
    assert outcome["contact_type"] == "dm_direct"
